"""Courriers PDF d'information des locataires (révision appliquée, gelée ou en attente d'indice).

Un PDF par (bail, échéance). Le texte vient de ``config/courrier_modele.json`` ;
les chiffres viennent du moteur ``calcul`` (donc identiques au classeur).
"""
from __future__ import annotations

import json
import logging
import shutil
from datetime import date, datetime, timedelta
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import calcul
from .baux import Bail, est_gel
from .insee import Observation, charger_config
from .workbook import R, Saisies, Societe, lire_classeur

log = logging.getLogger(__name__)
RACINE = Path(__file__).resolve().parent.parent
CHEMIN_MODELE = RACINE / "config" / "courrier_modele.json"

Cible = tuple[Bail, calcul.LigneRevision]


def charger_modele(chemin: Path = CHEMIN_MODELE) -> dict:
    with open(chemin, encoding="utf-8") as f:
        return json.load(f)


# ---------------------------------------------------------------------------
# Sélection des échéances à notifier
# ---------------------------------------------------------------------------
def selectionner(baux: list[Bail], observations: list[Observation], saisies: Saisies,
                 horizon_jours: int = 120, bail_id: str | None = None, tous: bool = False,
                 aujourdhui: date | None = None) -> list[Cible]:
    """Échéances à notifier : révision échue ou à venir dans ``horizon_jours``, dont le sort est
    connu (indice publié ou gel décidé) ou échue sans indice (courrier d'attente),
    non encore notifiée (« Courrier envoyé le » vide) et non appliquée, sauf ``tous``."""
    aujourdhui = aujourdhui or date.today()
    indices = calcul.table_indices(observations)
    limite = aujourdhui + timedelta(days=horizon_jours)
    cibles: list[Cible] = []
    for b in baux:
        if bail_id and b.id != bail_id:
            continue
        lignes = calcul.calculer(b, indices, horizon=b.date_fin or limite, aujourdhui=aujourdhui, decisions=saisies.decision)
        for l in lignes:
            cle = (b.id, l.echeance.numero)
            if l.echeance.date_revision > limite:
                continue
            if l.statut == calcul.STATUT_A_VENIR:
                continue  # ni indice ni décision : rien à dire
            if not tous and (saisies.courrier.get(cle) or str(saisies.applique.get(cle, "")).strip().lower() == "oui"):
                continue
            if l.loyer_precedent is None:
                continue  # chaîne interrompue en amont : impossible de chiffrer
            cibles.append((b, l))
    return cibles


# ---------------------------------------------------------------------------
# Génération
# ---------------------------------------------------------------------------
def _eur(x: float | None) -> str:
    if x is None:
        return "—"
    s = f"{x:,.2f}".replace(",", "\u00a0").replace(".", ",")
    return f"{s} €"


def _pct(x: float | None) -> str:
    return "—" if x is None else f"{x * 100:+.2f} %".replace(".", ",")


def _periodicite_adj(p: str) -> str:
    return {"Mensuelle": "mensuel", "Trimestrielle": "trimestriel", "Semestrielle": "semestriel", "Annuelle": "annuel"}.get(p, p.lower())


def valeurs_courrier(societe: Societe, bail: Bail, l: calcul.LigneRevision) -> dict[str, str]:
    cfg = charger_config()["series"].get(bail.indice, {})
    n = bail.echeances_par_an
    nouveau = l.loyer_retenu if not est_gel(l.decision) else l.loyer_precedent
    return {
        "locataire": bail.locataire, "local": bail.local, "indice": bail.indice,
        "indice_libelle": cfg.get("libelle", bail.indice).split(" - ")[0], "idbank": cfg.get("idbank", ""),
        "trimestre_reference": l.echeance.trimestre_reference, "valeur_reference": f"{l.indice_reference:g}".replace(".", ",") if l.indice_reference else "non publié",
        "trimestre_precedent": l.trimestre_precedent_effectif, "valeur_precedente": f"{l.indice_precedent:g}".replace(".", ",") if l.indice_precedent else "—",
        "coefficient": f"{l.coefficient:.6f}".replace(".", ",") if l.coefficient else "—",
        "loyer_precedent_annuel": _eur(l.loyer_precedent), "loyer_nouveau_annuel": _eur(nouveau),
        "loyer_precedent_echeance": _eur(l.loyer_precedent / n if l.loyer_precedent is not None else None),
        "loyer_nouveau_echeance": _eur(nouveau / n if nouveau is not None else None),
        "periodicite": _periodicite_adj(bail.periodicite_facturation), "variation_pct": _pct(l.variation_pct),
        "date_revision": f"{l.echeance.date_revision:%d/%m/%Y}", "date_effet_bail": f"{bail.date_effet:%d/%m/%Y}",
        "plafond": f"{bail.plafond_annuel_pct:g} % par an".replace(".", ",") if bail.plafond_annuel_pct is not None else "",
        "bailleur": societe.nom, "signataire": societe.signataire or societe.nom, "qualite": societe.qualite_signataire,
    }


def _styles():
    ss = getSampleStyleSheet()
    base = ParagraphStyle("base", parent=ss["Normal"], fontName="Helvetica", fontSize=10.5, leading=14, alignment=TA_JUSTIFY, spaceAfter=7)
    return {
        "base": base,
        "droite": ParagraphStyle("droite", parent=base, alignment=TA_RIGHT),
        "entete": ParagraphStyle("entete", parent=base, fontName="Helvetica-Bold", fontSize=11, alignment=0, spaceAfter=0),
        "petit": ParagraphStyle("petit", parent=base, fontSize=8, leading=10, textColor=colors.grey),
        "objet": ParagraphStyle("objet", parent=base, fontName="Helvetica-Bold", spaceBefore=6, spaceAfter=10, alignment=0),
        "demo": ParagraphStyle("demo", parent=base, fontName="Helvetica-Bold", textColor=colors.red, alignment=1),
    }


def _nl(texte: str) -> str:
    return texte.replace("\n", "<br/>")


def generer_un(societe: Societe, bail: Bail, l: calcul.LigneRevision, chemin: Path, modele: dict | None = None,
               aujourdhui: date | None = None) -> Path:
    modele = modele or charger_modele()
    aujourdhui = aujourdhui or date.today()
    v = valeurs_courrier(societe, bail, l)
    st = _styles()
    gel = est_gel(l.decision)
    attente = (not gel) and l.loyer_retenu is None

    doc = SimpleDocTemplate(str(chemin), pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=18 * mm, bottomMargin=22 * mm,
                            title=f"Révision de loyer – {bail.local}", author=societe.nom)
    mention = modele["mention_bas_de_page"].format(**v)

    def _pied(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.grey)
        largeur = A4[0] - 44 * mm
        texte = mention
        while canvas.stringWidth(texte, "Helvetica", 7.5) > largeur and len(texte) > 10:
            texte = texte[:-4] + "…"
        canvas.drawString(22 * mm, 12 * mm, texte)
        if _doc.page > 1:
            canvas.drawRightString(A4[0] - 22 * mm, 16 * mm, f"Page {_doc.page}")
        canvas.restoreState()
    el = []
    if societe.demo:
        el += [Paragraph("DOCUMENT DE DÉMONSTRATION – VALEURS FICTIVES", st["demo"]), Spacer(1, 4 * mm)]
    el += [Paragraph(f"<b>{societe.nom}</b>", st["entete"])]
    if societe.adresse:
        el.append(Paragraph(_nl(societe.adresse), st["base"]))
    if societe.siren:
        el.append(Paragraph(f"SIREN {societe.siren}", st["petit"]))
    el += [Spacer(1, 8 * mm),
           Paragraph(f"<b>{bail.locataire}</b><br/>{_nl(bail.adresse_courrier)}", st["droite"]),
           Spacer(1, 6 * mm),
           Paragraph(f"{societe.ville_signature + ', ' if societe.ville_signature else ''}le {aujourdhui:%d/%m/%Y}", st["droite"]),
           Spacer(1, 4 * mm)]
    objet = modele["objet_gel" if gel else "objet_attente" if attente else "objet_application"].format(**v)
    el += [Paragraph(f"Objet : {objet}", st["objet"]),
           Paragraph(f"Bail du {v['date_effet_bail']} – locaux : {bail.local}", st["petit"]),
           Spacer(1, 3 * mm), Paragraph(modele["intro"], st["base"])]
    if gel:
        paras = list(modele["gel"])
        if l.decision.endswith("rattrapage possible") or bail.methode == "Base fixe":
            paras.append(modele["gel_rattrapage"])
    elif attente:
        paras = list(modele["attente"])
    else:
        paras = list(modele["application"])
        if bail.plafond_annuel_pct is not None and l.loyer_brut is not None and l.loyer_retenu is not None and l.loyer_retenu < l.loyer_brut:
            paras.append(modele["application_plafond"])
    for ptxt in paras:
        el.append(Paragraph(ptxt.format(**v), st["base"]))

    if not attente:
        lignes_tab = [["", "Indice / loyer"],
                      [f"Indice {v['indice']} de référence ({v['trimestre_reference']})", v["valeur_reference"]],
                      [f"Indice {v['indice']} précédent ({v['trimestre_precedent']})", v["valeur_precedente"]],
                      ["Coefficient", v["coefficient"]],
                      ["Loyer annuel HT avant révision", v["loyer_precedent_annuel"]],
                      ["Loyer annuel HT " + ("maintenu" if gel else "révisé"), v["loyer_nouveau_annuel"]],
                      [f"Loyer {v['periodicite']} HT", v["loyer_nouveau_echeance"]]]
        if not gel:
            lignes_tab.append(["Variation", v["variation_pct"]])
        t = Table(lignes_tab, colWidths=[105 * mm, 45 * mm], hAlign="LEFT")
        t.setStyle(TableStyle([
            ("FONT", (0, 0), (-1, -1), "Helvetica", 9.5), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 9.5),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E1F2")), ("ALIGN", (1, 0), (1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
        ]))
        el += [Spacer(1, 2 * mm), t, Spacer(1, 5 * mm)]
    el += [Paragraph(modele["cloture"], st["base"]), Spacer(1, 8 * mm),
           Paragraph(f"{v['signataire']}<br/>{v['qualite']}, {societe.nom}", st["droite"])]
    chemin.parent.mkdir(parents=True, exist_ok=True)
    doc.build(el, onFirstPage=_pied, onLaterPages=_pied)
    return chemin


def generer(societe: Societe, cibles: list[Cible], dossier: Path, aujourdhui: date | None = None) -> list[Path]:
    modele = charger_modele()
    chemins = []
    for b, l in cibles:
        nature = "gel" if est_gel(l.decision) else "attente_indice" if l.loyer_retenu is None else "revision"
        nom = f"{_slug(b.id)}_{l.echeance.date_revision:%Y-%m-%d}_{nature}_{_slug(b.locataire)[:30]}.pdf"
        chemins.append(generer_un(societe, b, l, dossier / nom, modele, aujourdhui))
    return chemins


def marquer_envoyes(chemin_classeur: Path, cibles: list[Cible], aujourdhui: date | None = None) -> int:
    """Inscrit la date du jour dans « Courrier envoyé le » pour les échéances notifiées (avec sauvegarde)."""
    from openpyxl import load_workbook
    aujourdhui = aujourdhui or date.today()
    shutil.copy2(chemin_classeur, chemin_classeur.with_suffix(f".{datetime.now():%Y%m%d-%H%M%S}.bak.xlsx"))
    wb = load_workbook(chemin_classeur)
    ws = wb["Révisions"]
    cles = {(b.id, l.echeance.numero) for b, l in cibles}
    n = 0
    for r in range(2, ws.max_row + 1):
        cle = (str(ws[f"{R['ID bail']}{r}"].value), ws[f"{R['N°']}{r}"].value)
        if cle[1] is not None and (cle[0], int(cle[1])) in cles:
            c = ws[f"{R['Courrier envoyé le']}{r}"]
            c.value = aujourdhui
            c.number_format = "DD/MM/YYYY"
            n += 1
    wb.calculation.fullCalcOnLoad = True
    wb.save(chemin_classeur)
    return n


def _slug(s: str) -> str:
    return "".join(c if c.isalnum() else "_" for c in s).strip("_")
