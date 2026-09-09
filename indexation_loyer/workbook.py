"""Génération et rafraîchissement du classeur Excel de suivi d'une société.

Feuilles :
  Lisez-moi   : mode d'emploi, derniers indices connus, avertissements
  Société     : paramètres (raison sociale, SIREN, TVA...)
  Baux        : registre saisi par le cabinet (cellules jaunes = saisie)
  Indices     : cache INSEE recopié (une ligne par série × trimestre)
  Révisions   : calendrier des révisions, une ligne par bail × échéance,
                calculs par formules vivantes (SUMIFS sur « Indices »)
  Alertes     : synthèse par bail : loyer actuel, prochaine révision, action
  Pennylane   : préparation des abonnements de facturation (aperçu JSON)

Toutes les formules sont écrites en syntaxe « fichier » (anglais, séparateur
virgule) : Excel et LibreOffice les affichent localisés.
"""
from __future__ import annotations

import logging
import shutil
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from . import __version__
from .baux import INDICES, METHODES, PERIODICITES_FACTURATION, TYPES_BAIL, Bail
from .insee import Observation, charger_config, dernieres_valeurs

log = logging.getLogger(__name__)

# --- styles -----------------------------------------------------------------
BLEU = "1F3864"
FILL_ENTETE = PatternFill("solid", fgColor=BLEU)
FILL_SAISIE = PatternFill("solid", fgColor="FFF2CC")
FILL_DEMO = PatternFill("solid", fgColor="F8CBAD")
FILL_SECTION = PatternFill("solid", fgColor="D9E1F2")
FONT_ENTETE = Font(bold=True, color="FFFFFF")
FONT_TITRE = Font(bold=True, size=16, color=BLEU)
FONT_GRAS = Font(bold=True)
FONT_GRIS = Font(italic=True, color="7F7F7F")
BORDURE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
FMT_EUR = '#,##0.00\\ "€"'
FMT_PCT = "0.00%"
FMT_DATE = "DD/MM/YYYY"
FMT_INDICE = "0.00"

MAX_LIGNES_BAUX = 200      # plage couverte par les validations de données


@dataclass
class Societe:
    nom: str
    siren: str = ""
    forme: str = "SCI"
    regime_tva: str = "Loyers soumis à TVA sur option (20 %)"
    pennylane_company: str = ""
    contact: str = ""
    demo: bool = False


@dataclass
class Saisies:
    """Colonnes saisies à la main dans « Révisions », à préserver au rafraîchissement."""
    applique: dict[tuple[str, int], str] = field(default_factory=dict)
    date_application: dict[tuple[str, int], object] = field(default_factory=dict)
    commentaire: dict[tuple[str, int], str] = field(default_factory=dict)


# --- colonnes de la feuille Baux -----------------------------------------------
COLS_BAUX = [
    ("ID bail", 12), ("Local (désignation / adresse)", 34), ("Locataire", 26), ("Type de bail", 14),
    ("Date de prise d'effet", 14), ("Durée (ans)", 9), ("Date de fin", 14),
    ("Loyer initial annuel HT", 16), ("Charges annuelles HT", 14), ("TVA (%)", 8),
    ("Périodicité de facturation", 16), ("Indice", 8), ("Trimestre indice de base", 14),
    ("Périodicité de révision (ans)", 12), ("Méthode", 11), ("Plafond annuel (%) – optionnel", 14),
    ("Pennylane customer_id", 14), ("Pennylane product_id", 14), ("Notes", 30), ("Contrôles", 40),
]
# lettres utiles
B = {nom: get_column_letter(i + 1) for i, (nom, _) in enumerate(COLS_BAUX)}
B_ID, B_LOCAL, B_LOCATAIRE, B_TYPE = B["ID bail"], B["Local (désignation / adresse)"], B["Locataire"], B["Type de bail"]
B_EFFET, B_DUREE, B_FIN = B["Date de prise d'effet"], B["Durée (ans)"], B["Date de fin"]
B_LOYER, B_CHARGES, B_TVA = B["Loyer initial annuel HT"], B["Charges annuelles HT"], B["TVA (%)"]
B_PERFACT, B_INDICE, B_TBASE = B["Périodicité de facturation"], B["Indice"], B["Trimestre indice de base"]
B_PERREV, B_METHODE, B_PLAFOND = B["Périodicité de révision (ans)"], B["Méthode"], B["Plafond annuel (%) – optionnel"]
B_CUST, B_PROD, B_NOTES, B_CTRL = B["Pennylane customer_id"], B["Pennylane product_id"], B["Notes"], B["Contrôles"]

COLS_INDICES = [("Série", 8), ("Période", 10), ("Valeur", 10), ("Statut obs.", 9), ("Dernière MAJ INSEE", 14),
                ("Source", 12), ("Date extraction", 12), ("Variation annuelle", 11)]

COLS_REV = [
    ("ID bail", 10), ("Local", 28), ("N°", 5), ("Date de révision", 13), ("Indice", 7),
    ("Trim. référence", 11), ("Valeur indice réf.", 11), ("Trim. précédent", 11), ("Valeur indice préc.", 11),
    ("Coefficient", 11), ("Base de calcul (annuel HT)", 15), ("Loyer révisé brut", 14),
    ("Loyer retenu (annuel HT)", 15), ("Loyer précédent (annuel HT)", 15), ("Variation", 9),
    ("Loyer mensuel HT", 13), ("Loyer / échéance HT", 13), ("Charges / échéance HT", 13),
    ("TVA / échéance", 12), ("Total TTC / échéance", 13), ("Statut", 20),
    ("Appliqué ? (Oui/Non)", 11), ("Date d'application", 13), ("Commentaire", 30), ("Clé", 12),
]
R = {nom: get_column_letter(i + 1) for i, (nom, _) in enumerate(COLS_REV)}
COL_DATE_APPLI = R["Date d'application"]

COLS_ALERTES = [
    ("ID bail", 10), ("Local", 28), ("Locataire", 22), ("Indice", 7), ("Loyer initial annuel HT", 14),
    ("N° dernière révision calculable", 12), ("Loyer actuel (annuel HT)", 15), ("Loyer actuel mensuel HT", 13),
    ("Prochaine révision", 13), ("Trimestre attendu", 11), ("Indice publié ?", 10),
    ("Révisions calculables non appliquées", 14), ("Action", 34),
]
A = {nom: get_column_letter(i + 1) for i, (nom, _) in enumerate(COLS_ALERTES)}

COLS_PL = [
    ("ID bail", 10), ("Locataire", 22), ("customer_id", 12), ("product_id", 12), ("Libellé de ligne", 34),
    ("Périodicité", 12), ("recurrence.type", 12), ("Prix unitaire HT / échéance", 14),
    ("Charges HT / échéance", 13), ("TVA (%)", 8), ("vat_rate (code Pennylane)", 12),
    ("Date de début (1er du mois suivant)", 14), ("Aperçu du corps JSON (POST /billing_subscriptions)", 90),
]
P = {nom: get_column_letter(i + 1) for i, (nom, _) in enumerate(COLS_PL)}


# =============================================================================
# Construction
# =============================================================================
def construire(societe: Societe, baux: list[Bail], observations: list[Observation],
               chemin: Path, saisies: Saisies | None = None, aujourdhui: date | None = None) -> Path:
    aujourdhui = aujourdhui or date.today()
    saisies = saisies or Saisies()
    wb = Workbook()
    wb.remove(wb.active)
    _feuille_lisezmoi(wb, societe, observations, aujourdhui)
    _feuille_societe(wb, societe, aujourdhui)
    _feuille_baux(wb, baux)
    _feuille_indices(wb, observations)
    _feuille_revisions(wb, baux, saisies, aujourdhui)
    _feuille_alertes(wb, baux)
    _feuille_pennylane(wb, baux)
    wb.calculation.fullCalcOnLoad = True
    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb.save(chemin)
    log.info("Classeur écrit : %s", chemin)
    return chemin


def _entetes(ws: Worksheet, colonnes: list[tuple[str, int]], ligne: int = 1) -> None:
    for i, (nom, largeur) in enumerate(colonnes, start=1):
        c = ws.cell(row=ligne, column=i, value=nom)
        c.fill, c.font, c.border = FILL_ENTETE, FONT_ENTETE, BORDURE
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.row_dimensions[ligne].height = 42
    ws.freeze_panes = ws.cell(row=ligne + 1, column=1)


def _lookup(colonne_baux: str, ligne: int, feuille_col_id: str = "$A") -> str:
    """INDEX/MATCH de la colonne `colonne_baux` de Baux pour l'ID bail en colonne A de la ligne."""
    return f"INDEX(Baux!${colonne_baux}:${colonne_baux},MATCH({feuille_col_id}{ligne},Baux!$A:$A,0))"


# --- Lisez-moi ---------------------------------------------------------------
def _feuille_lisezmoi(wb: Workbook, societe: Societe, observations: list[Observation], aujourdhui: date) -> None:
    ws = wb.create_sheet("Lisez-moi")
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    for col in "CDEFG":
        ws.column_dimensions[col].width = 22
    ws["B2"] = f"Suivi des réindexations de loyers – {societe.nom}"
    ws["B2"].font = FONT_TITRE
    ws["B3"] = f"Généré le {aujourdhui:%d/%m/%Y} – indexation_loyer v{__version__}"
    ws["B3"].font = FONT_GRIS
    l = 5
    if societe.demo:
        ws.cell(row=l, column=2, value="⚠ CLASSEUR DE DÉMONSTRATION : tous les indices et tous les baux sont FICTIFS. "
                                        "Ne pas utiliser pour un calcul réel.").font = Font(bold=True, color="C00000")
        for c in range(2, 8):
            ws.cell(row=l, column=c).fill = FILL_DEMO
        l += 2

    ws.cell(row=l, column=2, value="Derniers indices connus dans ce classeur").font = FONT_GRAS
    l += 1
    for i, nom in enumerate(["Série", "Dernier trimestre", "Valeur", "Statut", "MAJ INSEE", "Source"], start=2):
        c = ws.cell(row=l, column=i, value=nom)
        c.fill, c.font = FILL_SECTION, FONT_GRAS
    config = charger_config()
    derniers = dernieres_valeurs(observations)
    for code in config["series"]:
        l += 1
        o = derniers.get(code)
        ws.cell(row=l, column=2, value=f"{code} – {config['series'][code]['libelle']}")
        if o:
            ws.cell(row=l, column=3, value=o.periode)
            ws.cell(row=l, column=4, value=o.valeur).number_format = FMT_INDICE
            ws.cell(row=l, column=5, value={"A": "Définitif", "P": "Provisoire"}.get(o.statut_obs, o.statut_obs))
            ws.cell(row=l, column=6, value=o.date_maj_insee)
            ws.cell(row=l, column=7, value=o.source)
        else:
            ws.cell(row=l, column=3, value="aucune valeur – lancer « fetch-indices »").font = FONT_GRIS
    l += 2
    ws.cell(row=l, column=2, value="Mode d'emploi").font = FONT_GRAS
    for texte in [
        "1. Feuille « Baux » : une ligne par local / bail. Seules les cellules jaunes se saisissent. La colonne Contrôles signale les incohérences.",
        "2. Feuille « Indices » : alimentée par la commande  python -m indexation_loyer fetch-indices  puis  refresh <société>. Ne pas saisir ici sauf ligne de source MANUEL.",
        "3. Feuille « Révisions » : calendrier calculé. Statut ✔ Calculable = l'indice est publié et la date de révision est passée : la facturation doit être ajustée.",
        "   Renseigner « Appliqué ? » et la date d'application une fois la révision facturée : ces colonnes sont conservées au rafraîchissement.",
        "4. Feuille « Alertes » : vue de pilotage par bail (loyer actuel, prochaine échéance, action).",
        "5. Feuille « Pennylane » : corps de requête prêt pour POST /billing_subscriptions (mode aperçu ; l'envoi se fait par la commande  pennylane <société>).",
        "Trimestres : notation AAAA-Tn (2024-T2 = 2e trimestre 2024), équivalente au 2024-Q2 de l'INSEE.",
        "Convention par défaut du trimestre de base : dernier indice publié à la date de prise d'effet (T-2). À caler sur la clause du bail.",
    ]:
        l += 1
        ws.cell(row=l, column=2, value=texte)
    l += 2
    ws.cell(row=l, column=2, value="Points de vigilance juridiques (à confirmer par le juriste du dossier)").font = FONT_GRAS
    for texte in [
        "• Depuis la loi Pinel (n° 2014-626 du 18 juin 2014), l'ILC et l'ILAT sont les indices de référence des baux commerciaux (art. L145-34 et L145-38 C. com.) ; l'ICC subsiste dans les clauses des baux antérieurs.",
        "• Plafonnement légal de la variation de l'ILC à 3,5 % pour les PME (loi n° 2022-1158 du 16 août 2022, art. 14, puis prolongation) : période et champ à vérifier bail par bail ; utiliser la colonne « Plafond annuel (%) ».",
        "• Une clause d'indexation ne jouant qu'à la hausse est réputée non écrite (art. L112-1 C. mon. fin. et jurisprudence Cass. 3e civ.) : le classeur applique la baisse si l'indice recule.",
        "• Révision légale triennale (L145-38) et déplafonnement : hors périmètre de la version actuelle, qui traite la clause d'échelle mobile contractuelle.",
        "• Le classeur ne remplace pas la lecture du bail : trimestre de base, périodicité et méthode (chaînée / base fixe) doivent être relevés sur l'acte.",
    ]:
        l += 1
        ws.cell(row=l, column=2, value=texte)
    l += 2
    ws.cell(row=l, column=2, value="Légende : fond jaune = saisie ; autres cellules = formules ; ne pas insérer de colonnes dans Baux, Indices, Révisions.").font = FONT_GRIS


# --- Société -----------------------------------------------------------------
def _feuille_societe(wb: Workbook, s: Societe, aujourdhui: date) -> None:
    ws = wb.create_sheet("Société")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 50
    lignes = [("Raison sociale", s.nom), ("SIREN", s.siren), ("Forme", s.forme),
              ("Régime TVA des loyers", s.regime_tva), ("Identifiant société Pennylane", s.pennylane_company),
              ("Contact cabinet", s.contact), ("Date de génération", aujourdhui), ("Mode démonstration", "OUI" if s.demo else "non")]
    for i, (k, v) in enumerate(lignes, start=1):
        ws.cell(row=i, column=1, value=k).font = FONT_GRAS
        c = ws.cell(row=i, column=2, value=v)
        c.fill = FILL_SAISIE
        if isinstance(v, date):
            c.number_format = FMT_DATE


# --- Baux --------------------------------------------------------------------
def _feuille_baux(wb: Workbook, baux: list[Bail]) -> None:
    ws = wb.create_sheet("Baux")
    _entetes(ws, COLS_BAUX)
    for i, b in enumerate(baux, start=2):
        valeurs = {
            B_ID: b.id, B_LOCAL: b.local, B_LOCATAIRE: b.locataire, B_TYPE: b.type_bail,
            B_EFFET: b.date_effet, B_DUREE: b.duree_ans, B_LOYER: b.loyer_initial_annuel_ht,
            B_CHARGES: b.charges_annuelles_ht, B_TVA: b.tva_pct, B_PERFACT: b.periodicite_facturation,
            B_INDICE: b.indice, B_TBASE: b.trimestre_base, B_PERREV: b.periodicite_revision_ans,
            B_METHODE: b.methode, B_PLAFOND: b.plafond_annuel_pct, B_CUST: b.pennylane_customer_id or None,
            B_PROD: b.pennylane_product_id or None, B_NOTES: b.notes or None,
        }
        for col, v in valeurs.items():
            ws[f"{col}{i}"] = v
        _formules_ligne_baux(ws, i)
    for i in range(2, max(len(baux) + 2, 3)):
        for nom, _ in COLS_BAUX:
            col = B[nom]
            if col not in (B_FIN, B_CTRL):
                ws[f"{col}{i}"].fill = FILL_SAISIE
            ws[f"{col}{i}"].border = BORDURE
        ws[f"{B_EFFET}{i}"].number_format = FMT_DATE
        ws[f"{B_FIN}{i}"].number_format = FMT_DATE
        for col in (B_LOYER, B_CHARGES):
            ws[f"{col}{i}"].number_format = FMT_EUR
    if not baux:
        _formules_ligne_baux(ws, 2)

    # validations de données
    plage = lambda col: f"{col}2:{col}{MAX_LIGNES_BAUX}"
    for col, liste in [(B_TYPE, TYPES_BAIL), (B_PERFACT, tuple(PERIODICITES_FACTURATION)),
                       (B_INDICE, INDICES), (B_METHODE, METHODES), (B_PERREV, ("1", "2", "3"))]:
        dv = DataValidation(type="list", formula1='"' + ",".join(liste) + '"', allow_blank=True)
        dv.add(plage(col))
        ws.add_data_validation(dv)
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="DATE(1990,1,1)", allow_blank=True)
    dv_date.add(plage(B_EFFET))
    ws.add_data_validation(dv_date)
    ws.conditional_formatting.add(f"{B_CTRL}2:{B_CTRL}{MAX_LIGNES_BAUX}",
                                  FormulaRule(formula=[f'AND({B_CTRL}2<>"",{B_CTRL}2<>"OK")'],
                                              font=Font(color="C00000", bold=True)))


def _formules_ligne_baux(ws: Worksheet, i: int) -> None:
    ws[f"{B_FIN}{i}"] = f'=IF(OR({B_EFFET}{i}="",{B_DUREE}{i}=""),"",EDATE({B_EFFET}{i},12*{B_DUREE}{i}))'
    ws[f"{B_FIN}{i}"].number_format = FMT_DATE
    ws[f"{B_CTRL}{i}"] = (
        f'=IF({B_ID}{i}="","",TRIM('
        f'IF({B_EFFET}{i}="","Date d\'effet manquante. ","")'
        f'&IF(NOT(ISNUMBER({B_LOYER}{i})),"Loyer initial manquant. ","")'
        f'&IF(COUNTIF(Indices!$A:$A,{B_INDICE}{i})=0,"Aucune valeur d\'indice chargée pour cette série. ","")'
        f'&IF(SUMIFS(Indices!$C:$C,Indices!$A:$A,{B_INDICE}{i},Indices!$B:$B,{B_TBASE}{i})=0,"Indice de base "&{B_TBASE}{i}&" absent de la feuille Indices. ","")'
        f'&IF(LEN({B_TBASE}{i})<>7,"Trimestre de base au format AAAA-Tn. ","")'
        f'&IF(COUNTIF($A$2:$A${MAX_LIGNES_BAUX},{B_ID}{i})>1,"ID bail en doublon. ","")'
        f'&IF(AND({B_PLAFOND}{i}<>"",NOT(ISNUMBER({B_PLAFOND}{i}))),"Plafond non numérique. ","")'
        f'&IF(COUNTIF($A$2:$A${MAX_LIGNES_BAUX},{B_ID}{i})=1,"","")'
        f')&IF(AND({B_EFFET}{i}<>"",ISNUMBER({B_LOYER}{i}),SUMIFS(Indices!$C:$C,Indices!$A:$A,{B_INDICE}{i},Indices!$B:$B,{B_TBASE}{i})>0,LEN({B_TBASE}{i})=7,COUNTIF($A$2:$A${MAX_LIGNES_BAUX},{B_ID}{i})=1),"OK",""))'
    )


# --- Indices -----------------------------------------------------------------
def _feuille_indices(wb: Workbook, observations: list[Observation]) -> None:
    ws = wb.create_sheet("Indices")
    _entetes(ws, COLS_INDICES)
    tri = sorted(observations, key=lambda o: (o.serie, o.periode))
    for i, o in enumerate(tri, start=2):
        ws.cell(row=i, column=1, value=o.serie)
        ws.cell(row=i, column=2, value=o.periode)
        ws.cell(row=i, column=3, value=o.valeur).number_format = FMT_INDICE
        ws.cell(row=i, column=4, value=o.statut_obs)
        ws.cell(row=i, column=5, value=o.date_maj_insee)
        ws.cell(row=i, column=6, value=o.source)
        ws.cell(row=i, column=7, value=o.date_extraction)
        c = ws.cell(row=i, column=8, value=(
            f'=IFERROR(C{i}/SUMIFS($C:$C,$A:$A,A{i},$B:$B,TEXT(VALUE(LEFT(B{i},4))-1,"0")&RIGHT(B{i},3))-1,"")'))
        c.number_format = FMT_PCT
        if o.source in ("FICTIF_DEMO", "MANUEL"):
            for col in range(1, 8):
                ws.cell(row=i, column=col).fill = FILL_DEMO if o.source == "FICTIF_DEMO" else FILL_SAISIE
    ws.auto_filter.ref = f"A1:H{max(len(tri) + 1, 2)}"


# --- Révisions ---------------------------------------------------------------
def _feuille_revisions(wb: Workbook, baux: list[Bail], saisies: Saisies, aujourdhui: date) -> None:
    ws = wb.create_sheet("Révisions")
    _entetes(ws, COLS_REV)
    i = 2
    for b in baux:
        horizon = b.date_fin or _plus_annees(aujourdhui, 3)
        premiere = i
        for ech in b.calendrier(horizon):
            _ligne_revision(ws, i, b, ech.numero, premiere)
            cle = (b.id, ech.numero)
            ws[f"{R['Appliqué ? (Oui/Non)']}{i}"] = saisies.applique.get(cle)
            ws[f"{COL_DATE_APPLI}{i}"] = saisies.date_application.get(cle)
            ws[f"{R['Commentaire']}{i}"] = saisies.commentaire.get(cle)
            i += 1
    derniere = max(i - 1, 2)
    dv = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
    dv.add(f"{R['Appliqué ? (Oui/Non)']}2:{R['Appliqué ? (Oui/Non)']}{derniere + 500}")
    ws.add_data_validation(dv)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS_REV))}{derniere}"
    col_statut = R["Statut"]
    ws.conditional_formatting.add(f"{col_statut}2:{col_statut}{derniere + 500}",
                                  FormulaRule(formula=[f'LEFT({col_statut}2,1)="⚠"'], font=Font(color="C00000", bold=True),
                                              fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.conditional_formatting.add(f"{col_statut}2:{col_statut}{derniere + 500}",
                                  FormulaRule(formula=[f'AND(LEFT({col_statut}2,1)="✔",{R["Appliqué ? (Oui/Non)"]}2<>"Oui")'],
                                              font=Font(color="9C5700", bold=True), fill=PatternFill("solid", fgColor="FFEB9C")))
    ws.conditional_formatting.add(f"{col_statut}2:{col_statut}{derniere + 500}",
                                  FormulaRule(formula=[f'AND(LEFT({col_statut}2,1)="✔",{R["Appliqué ? (Oui/Non)"]}2="Oui")'],
                                              font=Font(color="006100"), fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.column_dimensions[R["Clé"]].hidden = True


def _ligne_revision(ws: Worksheet, i: int, b: Bail, numero: int, premiere: int) -> None:
    lk = lambda col: _lookup(col, i)
    base_t = lk(B_TBASE)
    p = lk(B_PERREV)
    methode = lk(B_METHODE)
    loyer0 = lk(B_LOYER)
    plafond = lk(B_PLAFOND)
    c = lambda nom: f"{R[nom]}{i}"
    prec = lambda nom: f"{R[nom]}{i - 1}"
    ech_par_an = (f'IF({lk(B_PERFACT)}="Mensuelle",12,IF({lk(B_PERFACT)}="Trimestrielle",4,'
                  f'IF({lk(B_PERFACT)}="Semestrielle",2,1)))')

    ws[c("ID bail")] = b.id
    ws[c("Local")] = f"={lk(B_LOCAL)}"
    ws[c("N°")] = numero
    ws[c("Date de révision")] = f"=EDATE({lk(B_EFFET)},12*{p}*{c('N°')})"
    ws[c("Indice")] = f"={lk(B_INDICE)}"
    ws[c("Trim. référence")] = f'=TEXT(VALUE(LEFT({base_t},4))+{c("N°")}*{p},"0")&RIGHT({base_t},3)'
    ws[c("Valeur indice réf.")] = (f'=IF(SUMIFS(Indices!$C:$C,Indices!$A:$A,{c("Indice")},Indices!$B:$B,{c("Trim. référence")})=0,"",'
                                   f'SUMIFS(Indices!$C:$C,Indices!$A:$A,{c("Indice")},Indices!$B:$B,{c("Trim. référence")}))')
    ws[c("Trim. précédent")] = (f'=IF({methode}="Base fixe",{base_t},'
                                f'TEXT(VALUE(LEFT({base_t},4))+({c("N°")}-1)*{p},"0")&RIGHT({base_t},3))')
    ws[c("Valeur indice préc.")] = (f'=IF(SUMIFS(Indices!$C:$C,Indices!$A:$A,{c("Indice")},Indices!$B:$B,{c("Trim. précédent")})=0,"",'
                                    f'SUMIFS(Indices!$C:$C,Indices!$A:$A,{c("Indice")},Indices!$B:$B,{c("Trim. précédent")}))')
    ws[c("Coefficient")] = f'=IF(OR({c("Valeur indice réf.")}="",{c("Valeur indice préc.")}=""),"",{c("Valeur indice réf.")}/{c("Valeur indice préc.")})'
    if numero == 1:
        ws[c("Base de calcul (annuel HT)")] = f"={loyer0}"
        ws[c("Loyer précédent (annuel HT)")] = f"={loyer0}"
    else:
        ws[c("Base de calcul (annuel HT)")] = f'=IF({methode}="Base fixe",{loyer0},{prec("Loyer retenu (annuel HT)")})'
        ws[c("Loyer précédent (annuel HT)")] = f'={prec("Loyer retenu (annuel HT)")}'
    ws[c("Loyer révisé brut")] = f'=IF(OR({c("Coefficient")}="",{c("Base de calcul (annuel HT)")}=""),"",ROUND({c("Base de calcul (annuel HT)")}*{c("Coefficient")},2))'
    annees_plafond = f'IF({methode}="Base fixe",{c("N°")}*{p},{p})'
    ws[c("Loyer retenu (annuel HT)")] = (f'=IF({c("Loyer révisé brut")}="","",IF({plafond}="",{c("Loyer révisé brut")},'
                                         f'MIN({c("Loyer révisé brut")},ROUND({c("Base de calcul (annuel HT)")}*(1+{plafond}/100)^{annees_plafond},2))))')
    ws[c("Variation")] = f'=IF(OR({c("Loyer retenu (annuel HT)")}="",{c("Loyer précédent (annuel HT)")}=""),"",{c("Loyer retenu (annuel HT)")}/{c("Loyer précédent (annuel HT)")}-1)'
    ws[c("Loyer mensuel HT")] = f'=IF({c("Loyer retenu (annuel HT)")}="","",{c("Loyer retenu (annuel HT)")}/12)'
    ws[c("Loyer / échéance HT")] = f'=IF({c("Loyer retenu (annuel HT)")}="","",{c("Loyer retenu (annuel HT)")}/{ech_par_an})'
    ws[c("Charges / échéance HT")] = f'=IF({lk(B_CHARGES)}="",0,{lk(B_CHARGES)})/{ech_par_an}'
    ws[c("TVA / échéance")] = f'=IF({c("Loyer / échéance HT")}="","",({c("Loyer / échéance HT")}+{c("Charges / échéance HT")})*{lk(B_TVA)}/100)'
    ws[c("Total TTC / échéance")] = f'=IF({c("Loyer / échéance HT")}="","",{c("Loyer / échéance HT")}+{c("Charges / échéance HT")}+{c("TVA / échéance")})'
    ws[c("Statut")] = (f'=IF({c("Valeur indice réf.")}="",IF({c("Date de révision")}<=TODAY(),"⚠ Indice attendu","À venir"),'
                       f'IF({c("Date de révision")}<=TODAY(),"✔ Calculable","Indice connu – à venir"))')
    ws[c("Clé")] = f'={c("ID bail")}&"#"&{c("N°")}'

    ws[c("Date de révision")].number_format = FMT_DATE
    ws[c("Date d'application")].number_format = FMT_DATE
    for nom in ("Valeur indice réf.", "Valeur indice préc."):
        ws[c(nom)].number_format = FMT_INDICE
    ws[c("Coefficient")].number_format = "0.000000"
    ws[c("Variation")].number_format = FMT_PCT
    for nom in ("Base de calcul (annuel HT)", "Loyer révisé brut", "Loyer retenu (annuel HT)", "Loyer précédent (annuel HT)",
                "Loyer mensuel HT", "Loyer / échéance HT", "Charges / échéance HT", "TVA / échéance", "Total TTC / échéance"):
        ws[c(nom)].number_format = FMT_EUR
    for nom in ("Appliqué ? (Oui/Non)", "Date d'application", "Commentaire"):
        ws[c(nom)].fill = FILL_SAISIE
    for nom, _ in COLS_REV:
        ws[c(nom)].border = BORDURE


# --- Alertes -----------------------------------------------------------------
def _feuille_alertes(wb: Workbook, baux: list[Bail]) -> None:
    ws = wb.create_sheet("Alertes")
    _entetes(ws, COLS_ALERTES)
    rv = lambda nom: f"Révisions!${R[nom]}:${R[nom]}"
    for i, b in enumerate(baux, start=2):
        c = lambda nom: f"{A[nom]}{i}"
        lk = lambda col: _lookup(col, i)
        ws[c("ID bail")] = b.id
        ws[c("Local")] = f"={lk(B_LOCAL)}"
        ws[c("Locataire")] = f"={lk(B_LOCATAIRE)}"
        ws[c("Indice")] = f"={lk(B_INDICE)}"
        ws[c("Loyer initial annuel HT")] = f"={lk(B_LOYER)}"
        ws[c("N° dernière révision calculable")] = f'=_xlfn.MAXIFS({rv("N°")},{rv("ID bail")},{c("ID bail")},{rv("Statut")},"✔ Calculable")'
        ws[c("Loyer actuel (annuel HT)")] = (f'=IF({c("N° dernière révision calculable")}=0,{c("Loyer initial annuel HT")},'
                                             f'SUMIFS({rv("Loyer retenu (annuel HT)")},{rv("ID bail")},{c("ID bail")},{rv("N°")},{c("N° dernière révision calculable")}))')
        ws[c("Loyer actuel mensuel HT")] = f'={c("Loyer actuel (annuel HT)")}/12'
        ws[c("Prochaine révision")] = (f'=IF(_xlfn.MINIFS({rv("Date de révision")},{rv("ID bail")},{c("ID bail")},{rv("Date de révision")},">"&TODAY())=0,"—",'
                                       f'_xlfn.MINIFS({rv("Date de révision")},{rv("ID bail")},{c("ID bail")},{rv("Date de révision")},">"&TODAY()))')
        ws[c("Trimestre attendu")] = (f'=IFERROR(INDEX({rv("Trim. référence")},MATCH({c("ID bail")}&"#"&({c("N° dernière révision calculable")}+1),{rv("Clé")},0)),"—")')
        ws[c("Indice publié ?")] = (f'=IF({c("Trimestre attendu")}="—","—",IF(SUMIFS(Indices!$C:$C,Indices!$A:$A,{c("Indice")},Indices!$B:$B,{c("Trimestre attendu")})>0,"Oui","Non"))')
        ws[c("Révisions calculables non appliquées")] = (f'=COUNTIFS({rv("ID bail")},{c("ID bail")},{rv("Statut")},"✔ Calculable",'
                                                         f'{rv("Appliqué ? (Oui/Non)")},"<>Oui")')
        ws[c("Action")] = (f'=IF({c("Révisions calculables non appliquées")}>0,"⚠ Facturer la révision ("&{c("Révisions calculables non appliquées")}&" en attente)",'
                           f'IF(AND({c("Prochaine révision")}<>"—",{c("Prochaine révision")}-TODAY()<=60),'
                           f'IF({c("Indice publié ?")}="Oui","Révision dans "&INT({c("Prochaine révision")}-TODAY())&" j – indice publié","Révision dans "&INT({c("Prochaine révision")}-TODAY())&" j – indice non publié"),"RAS"))')
        for nom in ("Loyer initial annuel HT", "Loyer actuel (annuel HT)", "Loyer actuel mensuel HT"):
            ws[c(nom)].number_format = FMT_EUR
        ws[c("Prochaine révision")].number_format = FMT_DATE
        for nom, _ in COLS_ALERTES:
            ws[c(nom)].border = BORDURE
    n = max(len(baux) + 1, 2)
    ws.conditional_formatting.add(f"{A['Action']}2:{A['Action']}{n}",
                                  FormulaRule(formula=[f'LEFT({A["Action"]}2,1)="⚠"'], font=Font(color="C00000", bold=True),
                                              fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.conditional_formatting.add(f"{A['Action']}2:{A['Action']}{n}",
                                  FormulaRule(formula=[f'LEFT({A["Action"]}2,8)="Révision"'], fill=PatternFill("solid", fgColor="FFEB9C")))


# --- Pennylane ---------------------------------------------------------------
def _feuille_pennylane(wb: Workbook, baux: list[Bail]) -> None:
    ws = wb.create_sheet("Pennylane")
    _entetes(ws, COLS_PL)
    for i, b in enumerate(baux, start=2):
        c = lambda nom: f"{P[nom]}{i}"
        lk = lambda col: _lookup(col, i)
        al = lambda nom: f"INDEX(Alertes!${A[nom]}:${A[nom]},MATCH({c('ID bail')},Alertes!$A:$A,0))"
        ech_par_an = (f'IF({c("Périodicité")}="Mensuelle",12,IF({c("Périodicité")}="Trimestrielle",4,'
                      f'IF({c("Périodicité")}="Semestrielle",2,1)))')
        ws[c("ID bail")] = b.id
        ws[c("Locataire")] = f"={lk(B_LOCATAIRE)}"
        ws[c("customer_id")] = f'=IF({lk(B_CUST)}="","À RENSEIGNER",{lk(B_CUST)})'
        ws[c("product_id")] = f'=IF({lk(B_PROD)}="","",{lk(B_PROD)})'
        ws[c("Libellé de ligne")] = f'="Loyer "&{lk(B_LOCAL)}&" – échéance "&LOWER({c("Périodicité")})'
        ws[c("Périodicité")] = f"={lk(B_PERFACT)}"
        ws[c("recurrence.type")] = (f'=IF({c("Périodicité")}="Mensuelle","monthly",IF({c("Périodicité")}="Trimestrielle","quarterly",'
                                    f'IF({c("Périodicité")}="Semestrielle","semiannual","yearly")))')
        ws[c("Prix unitaire HT / échéance")] = f'=ROUND({al("Loyer actuel (annuel HT)")}/{ech_par_an},2)'
        ws[c("Charges HT / échéance")] = f'=ROUND(IF({lk(B_CHARGES)}="",0,{lk(B_CHARGES)})/{ech_par_an},2)'
        ws[c("TVA (%)")] = f"={lk(B_TVA)}"
        ws[c("vat_rate (code Pennylane)")] = (f'=IF({c("TVA (%)")}=20,"FR_200",IF({c("TVA (%)")}=10,"FR_100",'
                                              f'IF({c("TVA (%)")}=5.5,"FR_55",IF({c("TVA (%)")}=2.1,"FR_21","exempt"))))')
        ws[c("Date de début (1er du mois suivant)")] = "=DATE(YEAR(TODAY()),MONTH(TODAY())+1,1)"
        d = c("Date de début (1er du mois suivant)")
        iso = f'YEAR({d})&"-"&TEXT(MONTH({d}),"00")&"-"&TEXT(DAY({d}),"00")'
        num = lambda ref: f'SUBSTITUTE(TEXT({ref},"0.00"),",",".")'
        ligne_charges = (f'IF({c("Charges HT / échéance")}>0,", {{""label"": ""Provision sur charges"", ""quantity"": 1, '
                         f'""raw_currency_unit_price"": "&{num(c("Charges HT / échéance"))}&", ""vat_rate"": """&{c("vat_rate (code Pennylane)")}&"""}}","")')
        ws[c("Aperçu du corps JSON (POST /billing_subscriptions)")] = (
            f'="{{""customer_id"": "&{c("customer_id")}&", ""start"": """&{iso}&""", ""currency"": ""EUR"", '
            f'""recurrence"": {{""type"": """&{c("recurrence.type")}&""", ""day_of_month"": 1}}, '
            f'""invoice_lines"": [{{""label"": """&{c("Libellé de ligne")}&""", ""quantity"": 1, '
            f'""raw_currency_unit_price"": "&{num(c("Prix unitaire HT / échéance"))}&", ""vat_rate"": """&{c("vat_rate (code Pennylane)")}&""""'
            f'&IF({c("product_id")}<>"",", ""product_id"": "&{c("product_id")},"")&"}}"&{ligne_charges}&"]}}"'
        )
        for nom in ("Prix unitaire HT / échéance", "Charges HT / échéance"):
            ws[c(nom)].number_format = FMT_EUR
        ws[c("Date de début (1er du mois suivant)")].number_format = FMT_DATE
        ws[c("Aperçu du corps JSON (POST /billing_subscriptions)")].alignment = Alignment(wrap_text=False)
        for nom, _ in COLS_PL:
            ws[c(nom)].border = BORDURE
    n = len(baux) + 3
    ws.cell(row=n, column=1, value="Schéma indicatif : les noms de champs (recurrence, invoice_lines, raw_currency_unit_price, vat_rate) "
                                    "sont à valider contre https://pennylane.readme.io/reference/postbillingsubscriptions avant tout envoi. "
                                    "Le mapping est modifiable dans config/pennylane_mapping.json.").font = FONT_GRIS


# =============================================================================
# Lecture d'un classeur existant (pour refresh / pennylane)
# =============================================================================
def lire_classeur(chemin: Path) -> tuple[Societe, list[Bail], Saisies]:
    wb = load_workbook(chemin, data_only=False)
    ws_s = wb["Société"]
    params = {ws_s.cell(row=r, column=1).value: ws_s.cell(row=r, column=2).value for r in range(1, 12)}
    societe = Societe(
        nom=str(params.get("Raison sociale") or chemin.stem), siren=str(params.get("SIREN") or ""),
        forme=str(params.get("Forme") or ""), regime_tva=str(params.get("Régime TVA des loyers") or ""),
        pennylane_company=str(params.get("Identifiant société Pennylane") or ""),
        contact=str(params.get("Contact cabinet") or ""), demo=str(params.get("Mode démonstration")).upper() == "OUI",
    )
    ws_b = wb["Baux"]
    entetes = [ws_b.cell(row=1, column=c).value for c in range(1, len(COLS_BAUX) + 1)]
    attendus = [n for n, _ in COLS_BAUX]
    if entetes != attendus:
        raise ValueError("La feuille Baux a été modifiée (colonnes déplacées ou renommées) : "
                         f"attendu {attendus}, trouvé {entetes}")
    baux: list[Bail] = []
    for r in range(2, ws_b.max_row + 1):
        v = lambda col: ws_b[f"{col}{r}"].value
        if v(B_ID) in (None, ""):
            continue
        baux.append(Bail(
            id=str(v(B_ID)).strip(), local=str(v(B_LOCAL) or ""), locataire=str(v(B_LOCATAIRE) or ""),
            type_bail=str(v(B_TYPE) or "Commercial"), date_effet=_en_date(v(B_EFFET)),
            duree_ans=int(v(B_DUREE)) if v(B_DUREE) not in (None, "") else None,
            loyer_initial_annuel_ht=float(v(B_LOYER) or 0), charges_annuelles_ht=float(v(B_CHARGES) or 0),
            tva_pct=float(v(B_TVA) if v(B_TVA) not in (None, "") else 20),
            periodicite_facturation=str(v(B_PERFACT) or "Mensuelle"), indice=str(v(B_INDICE) or "ILC"),
            trimestre_base=str(v(B_TBASE) or ""), periodicite_revision_ans=int(v(B_PERREV) or 1),
            methode=str(v(B_METHODE) or "Chaînée"),
            plafond_annuel_pct=float(v(B_PLAFOND)) if v(B_PLAFOND) not in (None, "") else None,
            pennylane_customer_id=str(v(B_CUST) or "").replace(".0", "") if v(B_CUST) is not None else "",
            pennylane_product_id=str(v(B_PROD) or "").replace(".0", "") if v(B_PROD) is not None else "",
            notes=str(v(B_NOTES) or ""),
        ))
    saisies = Saisies()
    if "Révisions" in wb.sheetnames:
        ws_r = wb["Révisions"]
        for r in range(2, ws_r.max_row + 1):
            bid, num = ws_r[f"{R['ID bail']}{r}"].value, ws_r[f"{R['N°']}{r}"].value
            if bid in (None, "") or num in (None, ""):
                continue
            cle = (str(bid), int(num))
            for attr, nom in (("applique", "Appliqué ? (Oui/Non)"), ("date_application", "Date d'application"),
                              ("commentaire", "Commentaire")):
                val = ws_r[f"{R[nom]}{r}"].value
                if val not in (None, ""):
                    getattr(saisies, attr)[cle] = val
    return societe, baux, saisies


def rafraichir(chemin: Path, observations: list[Observation], aujourdhui: date | None = None) -> Path:
    societe, baux, saisies = lire_classeur(chemin)
    erreurs = [(b.id, e) for b in baux for e in b.erreurs]
    for bid, e in erreurs:
        log.warning("Bail %s : %s", bid, e)
    sauvegarde = chemin.with_suffix(f".{datetime.now():%Y%m%d-%H%M%S}.bak.xlsx")
    shutil.copy2(chemin, sauvegarde)
    log.info("Sauvegarde : %s", sauvegarde)
    return construire(societe, baux, observations, chemin, saisies=saisies, aujourdhui=aujourdhui)


def _en_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    raise ValueError(f"Date illisible : {v!r}")


def _plus_annees(d: date, n: int) -> date:
    try:
        return d.replace(year=d.year + n)
    except ValueError:
        return d.replace(year=d.year + n, day=28)
