"""Tests du classeur : structure, lecture/rafraîchissement, et parité des formules
recalculées par LibreOffice avec le moteur Python (ignoré si soffice indisponible)."""
import os
import shutil
import subprocess
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from indexation_loyer import calcul, demo
from indexation_loyer.baux import Bail
from indexation_loyer.workbook import (COLS_BAUX, R, A, P, B_LOYER, Societe, construire, lire_classeur, rafraichir)

AUJOURDHUI = date(2026, 9, 9)


@pytest.fixture(scope="module")
def classeur_demo(tmp_path_factory) -> Path:
    dossier = tmp_path_factory.mktemp("demo")
    return demo.generer_demo(dossier, aujourdhui=AUJOURDHUI)


def test_structure(classeur_demo):
    wb = load_workbook(classeur_demo)
    assert wb.sheetnames == ["Lisez-moi", "Société", "Baux", "Indices", "Révisions", "Alertes", "Pennylane"]
    ws = wb["Baux"]
    assert [ws.cell(row=1, column=i + 1).value for i in range(len(COLS_BAUX))] == [n for n, _ in COLS_BAUX]
    ws_r = wb["Révisions"]
    # B01 : bail de 9 ans, révision annuelle -> 9 lignes ; B02 triennal -> 3 lignes
    ids = [ws_r[f"A{r}"].value for r in range(2, ws_r.max_row + 1)]
    assert ids.count("B01") == 9 and ids.count("B02") == 3
    assert ws_r[f"{R['Statut']}2"].value.startswith("=IF(")
    assert "_xlfn.MAXIFS" in wb["Alertes"][f"{A['N° dernière révision effective']}2"].value


def test_lire_puis_rafraichir_conserve_saisies(classeur_demo, tmp_path):
    copie = tmp_path / "copie.xlsx"
    shutil.copy(classeur_demo, copie)
    wb = load_workbook(copie)
    ws = wb["Révisions"]
    ws[f"{R['Appliqué ? (Oui/Non)']}2"] = "Oui"
    ws[f"{R['Commentaire']}2"] = "Facturé le 05/04/2022"
    wb["Baux"][f"{B_LOYER}2"] = 25000  # modification du loyer initial de B01
    wb.save(copie)

    societe, baux, saisies = lire_classeur(copie)
    assert societe.demo and len(baux) == 4
    assert baux[0].loyer_initial_annuel_ht == 25000
    assert saisies.applique[("B01", 1)] == "Oui"
    assert saisies.decision[("B01", 4)] == "Geler – sans rattrapage"      # saisie de la démo relue
    assert saisies.consigne[("B03", 6)].startswith("Gel d'un an")

    rafraichir(copie, demo.indices_fictifs(AUJOURDHUI), aujourdhui=AUJOURDHUI)
    wb2 = load_workbook(copie)
    ws2 = wb2["Révisions"]
    assert ws2[f"{R['Appliqué ? (Oui/Non)']}2"].value == "Oui"
    assert ws2[f"{R['Commentaire']}2"].value == "Facturé le 05/04/2022"
    assert ws2[f"{R['Décision']}5"].value == "Geler – sans rattrapage"   # B01 n°4 conservée
    assert wb2["Baux"][f"{B_LOYER}2"].value == 25000
    assert list(tmp_path.glob("copie.*.bak.xlsx"))  # sauvegarde créée


def test_lire_classeur_refuse_colonnes_modifiees(classeur_demo, tmp_path):
    copie = tmp_path / "casse.xlsx"
    shutil.copy(classeur_demo, copie)
    wb = load_workbook(copie)
    wb["Baux"].insert_cols(3)
    wb.save(copie)
    with pytest.raises(ValueError, match="feuille Baux"):
        lire_classeur(copie)


def _recalculer(chemin: Path, tmp: Path) -> Path | None:
    soffice = shutil.which("soffice")
    if not soffice:
        return None
    env = dict(os.environ, HOME=str(tmp / "home"))
    (tmp / "home").mkdir(exist_ok=True)
    r = subprocess.run([soffice, "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", str(tmp / "out"), str(chemin)],
                       capture_output=True, text=True, timeout=300, env=env)
    sortie = tmp / "out" / chemin.name
    return sortie if sortie.exists() else None


def test_parite_libreoffice_moteur_python(classeur_demo, tmp_path):
    recalc = _recalculer(classeur_demo, tmp_path)
    if recalc is None:
        pytest.skip("LibreOffice Calc indisponible : parité non vérifiée")
    wb = load_workbook(recalc, data_only=True)

    # 1. aucune erreur de formule dans tout le classeur
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                assert not (isinstance(v, str) and v.startswith(("#", "Err:"))), f"{ws.title}: {v}"

    # 2. Révisions : loyer retenu / statut identiques au moteur Python
    indices = calcul.table_indices(demo.indices_fictifs(AUJOURDHUI))
    decisions = demo.saisies_fictives().decision
    attendu = {}
    baux = demo.baux_fictifs()
    for b in baux:
        for l in calcul.calculer(b, indices, horizon=b.date_fin, aujourdhui=AUJOURDHUI, decisions=decisions):
            attendu[(b.id, l.echeance.numero)] = l
    ws = wb["Révisions"]
    n = 0
    for r in range(2, ws.max_row + 1):
        cle = (ws[f"A{r}"].value, ws[f"{R['N°']}{r}"].value)
        l = attendu[cle]
        retenu = ws[f"{R['Loyer retenu (annuel HT)']}{r}"].value
        assert (retenu in (None, "")) == (l.loyer_retenu is None), cle
        if l.loyer_retenu is not None:
            assert abs(retenu - l.loyer_retenu) < 0.005, cle
        assert ws[f"{R['Statut']}{r}"].value == l.statut, cle
        assert ws[f"{R['Trim. précédent']}{r}"].value == l.trimestre_precedent_effectif, cle
        assert ws[f"{R['Date de révision']}{r}"].value.date() == l.echeance.date_revision
        n += 1
    assert n == len(attendu) == 30

    # 3. Alertes : loyer actuel = dernier calculable
    ws_a = wb["Alertes"]
    for r in range(2, ws_a.max_row + 1):
        b = next(b for b in baux if b.id == ws_a[f"A{r}"].value)
        lignes = calcul.calculer(b, indices, horizon=b.date_fin, aujourdhui=AUJOURDHUI, decisions=decisions)
        assert abs(ws_a[f"{A['Loyer actuel (annuel HT)']}{r}"].value - calcul.loyer_actuel(b, lignes, AUJOURDHUI)) < 0.005
    assert ws_a[f"{A['Révisions gelées']}2"].value == 1                      # B01
    assert ws_a[f"{A['Consigne prochaine révision']}4"].value in (None, "")          # B03 : prochaine = n°8, sans consigne
    # B03 : n°7 est la dernière effective ; la consigne affichée en Révisions n°7 est celle du rattrapage
    assert ws[f"{R['Consigne (à faire)']}{[r for r in range(2, ws.max_row + 1) if ws[f'A{r}'].value == 'B03' and ws[f'{R[chr(78) + chr(176)]}{r}'].value == 7][0]}"].value.startswith("Appliquer le rattrapage")
    # B01 : la révision 2027-04-01 n'est pas encore publiée
    assert ws_a[f"{A['Indice publié ?']}2"].value == "Non"
    assert ws_a[f"{A['Trimestre attendu']}2"].value == "2026-T4"
    assert ws_a[f"{A['Action']}2"].value.startswith("⚠ Facturer")

    # 4. Pennylane : le JSON d'aperçu est valide et cohérent
    import json
    ws_p = wb["Pennylane"]
    corps = json.loads(ws_p[f"{P['Aperçu du corps JSON (POST /billing_subscriptions)']}2"].value)
    assert corps["customer_id"] == 100001 and corps["recurrence"]["type"] == "monthly"
    assert len(corps["invoice_lines"]) == 2  # loyer + charges
    assert abs(corps["invoice_lines"][0]["raw_currency_unit_price"] - calcul.loyer_actuel(baux[0], calcul.calculer(baux[0], indices, baux[0].date_fin, AUJOURDHUI, decisions), AUJOURDHUI) / 12) < 0.01
